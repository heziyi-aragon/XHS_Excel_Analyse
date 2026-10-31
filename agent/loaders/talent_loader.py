"""TalentLoader：读取蒲公英达人数据表（记录式 Excel，单行表头 + 数据行）。

表头匹配用「关键字包含」而不是全等：真实表头是
'平台图文参考价（元）'、'合作笔记曝光数中位值（近30天）' 这种带备注的长名字，
关键字匹配让列序调整、备注修改都不影响读取。
"""
from __future__ import annotations

from pathlib import Path

import openpyxl

from agent.domain import LoadReport, TalentProfile, _TALENT_NUMERIC, _TALENT_TEXT


class TalentLoader:
    def __init__(self, path: str | Path, sheet_name: str = "达人数据"):
        self.path = Path(path)
        self.sheet_name = sheet_name

    def load(self, report: LoadReport | None = None) -> list[TalentProfile]:
        report = report or LoadReport()
        if not self.path.exists():
            report.add(f"达人数据文件不存在：{self.path}")
            return []

        ws = openpyxl.load_workbook(self.path, data_only=True)[self.sheet_name]
        headers = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]

        # 给每一列找到它对应的字段名和解析函数；同一字段只绑定最先出现的列，
        # 并排除「CPV优于行业top50%」这类衍生对比列对 CPV 原值的覆盖。
        col_map: dict[int, tuple[str, object]] = {}
        bound_fields: set[str] = set()
        for idx, header in enumerate(headers):
            if not header:
                continue
            for key, (field, parser) in _TALENT_NUMERIC.items():
                if field in bound_fields:
                    continue
                if key == "CPV" and "优于" in header:
                    continue
                if key in header:
                    col_map[idx] = (field, parser)
                    bound_fields.add(field)
                    break
            else:
                for key, field in _TALENT_TEXT.items():
                    if field in bound_fields:
                        continue
                    if key in header:
                        col_map[idx] = (field, None)
                        bound_fields.add(field)
                        break
                else:
                    continue  # 认识不了的列跳过，不报错

        talents: list[TalentProfile] = []
        for row in ws.iter_rows(min_row=2):
            if row[0].value is None:  # 空行
                continue
            talent = TalentProfile()
            for idx, (field, parser) in col_map.items():
                value = row[idx].value if idx < len(row) else None
                if parser is not None:
                    setattr(talent, field, parser(value))
                else:
                    setattr(talent, field, str(value).strip() if value is not None else None)
            # 表里认得、但这一行恰好没有的文本字段补成 None，保持对象形状一致
            for field in _TALENT_TEXT.values():
                if getattr(talent, field) == "":
                    setattr(talent, field, None)

            if not talent.nickname:
                report.add(f"第 {row[0].row} 行达人缺少昵称，该行已跳过")
                continue
            if talent.price_image is None and talent.price_video is None:
                report.add(f"达人「{talent.nickname}」图文与视频参考价均缺失")
            talents.append(talent)

        if len(talents) < 5:
            report.add(f"候选达人仅 {len(talents)} 人（少于 5 人），后续筛选可能不足")
        return talents
