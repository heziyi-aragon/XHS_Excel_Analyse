import pandas as pd
from openpyxl import load_workbook

class ColumnDivider:
    def __init__(self, file_name):
        self.input_path = f"static/{file_name}"
        self.output_path = f"static/{file_name.replace('.xlsx', '_处理后.xlsx')}"
        self.df = None
        self.first_row = None  # 保存第一行原始内容

    def read_excel(self):
        wb = load_workbook(self.input_path, data_only=False)
        ws = wb.active
        self.first_row = [cell.value for cell in ws[1]]  # 获取第一行所有单元格值
        wb.close()
        
        self.df = pd.read_excel(self.input_path, header=1)

    def add_recommend_column(self):
        def judge_recommend(val):
            if pd.isna(val):
                return None
            val=val.replace('，',',')
            prefix = val.split(',')[0].strip()
            if prefix == 'Y':
                return "TRUE"
            elif 'Y' in prefix:
                return "TRUE"
            else:
                return "FALSE"
            
        
        self.df['是否推荐'] = self.df['Social是否推荐及理由'].apply(judge_recommend)

    def add_reason_column(self):
        def extract_reason(val):
            if pd.isna(val):
                return None
            val=val.replace(",","，")
            if val[0] not in ['Y', 'N', 'TBD']:
                return val
            elif val[1]!="，":
                return val
            parts = val.split('，', 1)
            return parts[1].strip() if len(parts) > 1 else ''
        
        self.df['原因'] = self.df['Social是否推荐及理由'].apply(extract_reason)

    def save_excel(self):
        self.df.to_excel(self.output_path, index=False, header=True, startrow=1)
        
        wb = load_workbook(self.output_path, data_only=False)
        ws = wb.active
        for col_idx, value in enumerate(self.first_row, 1):
            ws.cell(row=1, column=col_idx, value=value)
        wb.save(self.output_path)
        wb.close()
        return self.output_path
